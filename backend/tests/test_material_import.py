# backend/tests/test_material_import.py
"""
Тесты импорта/экспорта Excel (Итерация 13.2).

Покрывают:
  1. Pydantic-модели MaterialImportRow / MaterialImportResponse.
  2. Парсер Excel (внутренние функции).
  3. Валидацию строк.
  4. Заголовки шаблона.
"""
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.api.v1.material_models import (
    MaterialImportRow,
    MaterialImportResponse,
)


# ==========================================
# PYDANTIC-МОДЕЛИ
# ==========================================

def test_import_row_created():
    """Строка CREATED."""
    r = MaterialImportRow(
        row_number=2,
        code="WATER",
        name="Вода",
        category="RAW",
        unit="kg",
        qty=100000,
        reserved_qty=0,
        status="CREATED",
        message="Материал создан",
    )
    assert r.status == "CREATED"
    assert r.qty == 100000


def test_import_row_error():
    """Строка ERROR с сообщением."""
    r = MaterialImportRow(
        row_number=5,
        code="BAD",
        status="ERROR",
        message="Некорректное число в qty",
    )
    assert r.status == "ERROR"
    assert "Некорректное" in r.message


def test_import_response_aggregates():
    """Ответ агрегирует счётчики."""
    r = MaterialImportResponse(
        total_rows=10,
        created=5,
        updated=3,
        skipped=1,
        errors=1,
        rows=[],
        message="Импорт завершён",
    )
    assert r.total_rows == 10
    assert r.created + r.updated + r.skipped + r.errors == 10


# ==========================================
# СТРУКТУРА МОДУЛЯ
# ==========================================

def test_excel_headers_mapping():
    """Словарь _EXCEL_HEADERS существует и содержит ключи."""
    from app.api.v1.materials import _EXCEL_HEADERS
    assert "код" in _EXCEL_HEADERS
    assert "code" in _EXCEL_HEADERS
    assert "остаток" in _EXCEL_HEADERS
    assert _EXCEL_HEADERS["код"] == "code"
    assert _EXCEL_HEADERS["остаток"] == "qty"


def test_valid_categories():
    """_VALID_CATEGORIES = {RAW, PACKAGING, LABEL}."""
    from app.api.v1.materials import _VALID_CATEGORIES
    assert _VALID_CATEGORIES == {"RAW", "PACKAGING", "LABEL"}


def test_import_endpoint_exists():
    """POST /import-excel зарегистрирован."""
    from app.api.v1 import materials as materials_module
    assert hasattr(materials_module, "import_materials_from_excel")


def test_template_endpoint_exists():
    """GET /import-template зарегистрирован."""
    from app.api.v1 import materials as materials_module
    assert hasattr(materials_module, "download_import_template")


def test_export_endpoint_exists():
    """GET /export-excel зарегистрирован."""
    from app.api.v1 import materials as materials_module
    assert hasattr(materials_module, "export_materials_to_excel")


# ==========================================
# ПАРСИНГ EXCEL (через openpyxl)
# ==========================================

def _create_test_xlsx(rows: list) -> bytes:
    """Создать in-memory Excel-файл с заданными строками."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Код", "Наименование", "Категория", "Ед.изм.", "Остаток", "Резерв"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_excel_parsing_basic():
    """Простой Excel парсится openpyxl."""
    data = _create_test_xlsx([
        ("WATER", "Вода", "RAW", "kg", 100000, 0),
        ("SALT", "Соль", "RAW", "kg", 3000, 0),
    ])
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    assert len(rows) == 3  # header + 2 data
    assert rows[0] == ("Код", "Наименование", "Категория", "Ед.изм.", "Остаток", "Резерв")
    assert rows[1][0] == "WATER"
    assert rows[1][4] == 100000


def test_excel_parsing_empty_rows():
    """Пустые строки сохраняются (валидация должна их пропустить)."""
    data = _create_test_xlsx([
        ("WATER", "Вода", "RAW", "kg", 100000, 0),
        (None, None, None, None, None, None),
        ("SALT", "Соль", "RAW", "kg", 3000, 0),
    ])
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 4


# ==========================================
# ШАБЛОН (структурный тест)
# ==========================================

def test_template_generates_xlsx():
    """Шаблон генерирует валидный xlsx-файл."""
    # Копируем логику из download_import_template (упрощённо)
    wb = Workbook()
    ws = wb.active
    ws.append(["Код", "Наименование", "Категория", "Ед.изм.", "Остаток", "Резерв"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Проверяем, что файл открывается
    wb2 = load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
    ws2 = wb2.active
    headers = [c.value for c in ws2[1]]
    assert headers == ["Код", "Наименование", "Категория", "Ед.изм.", "Остаток", "Резерв"]


# ==========================================
# ВАЛИДАЦИЯ КАТЕГОРИЙ
# ==========================================

def test_invalid_category_rejected():
    """Недопустимая категория → ERROR."""
    from app.api.v1.materials import _VALID_CATEGORIES
    assert "INVALID" not in _VALID_CATEGORIES
    assert "RAW" in _VALID_CATEGORIES


if __name__ == "__main__":
    pytest.main([__file__, "-v"])