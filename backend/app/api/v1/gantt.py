# backend/app/api/v1/gantt.py
"""
API для диаграммы Ганта.

Итерация 1:
- Возвращает linked_equipment_id, linked_equipment_name, task_role

Итерация 5 (hotfix):
- Параметр version_id. Фильтрация по schedule_version_id.

Итерация 5 (hotfix #3):
- is_lab_blocked, lab_status, lab_block_reason

Итерация 7:
- cooling_mode (fast | slow) для операций охлаждения с деградацией.
"""

import io
from datetime import datetime
from uuid import UUID

from fastapi import APIRouter, HTTPException, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.dependencies import get_current_org_id, get_db_session
from .models import GanttTask, GanttResponse
from .schedule import _last_schedule_result

router = APIRouter(prefix="/api/v1/gantt", tags=["Диаграмма Ганта"])


async def _resolve_version_id(
        db: AsyncSession,
        org_id: UUID,
        version_id: UUID | None,
) -> UUID | None:
    """Определяет версию плана для отображения."""
    if version_id is not None:
        result = await db.execute(
            text("""
                SELECT id FROM schedule_version
                WHERE id = :version_id AND organization_id = :org_id
            """),
            {"version_id": version_id, "org_id": org_id},
        )
        if result.fetchone():
            return version_id
        raise HTTPException(
            status_code=404,
            detail=f"Версия плана {version_id} не найдена",
        )

    result = await db.execute(
        text("""
            SELECT id FROM schedule_version
            WHERE organization_id = :org_id AND is_active = TRUE
            ORDER BY created_at DESC
            LIMIT 1
        """),
        {"org_id": org_id},
    )
    row = result.fetchone()
    if row:
        return row.id

    result = await db.execute(
        text("""
            SELECT id FROM schedule_version
            WHERE organization_id = :org_id
            ORDER BY created_at DESC
            LIMIT 1
        """),
        {"org_id": org_id},
    )
    row = result.fetchone()
    return row.id if row else None


@router.get("/", response_model=GanttResponse)
async def get_gantt_data(
        version_id: UUID | None = Query(default=None, description="ID сохраненной версии плана. Если не задан — последняя активная."),
        org_id: UUID = Depends(get_current_org_id),
        db: AsyncSession = Depends(get_db_session),
):
    """
    Получение данных для диаграммы Ганта.

    Логика:
      - Если version_id передан явно — загружаем из БД по нему.
      - Если не передан и есть сохранённые версии — берём последнюю активную.
      - Иначе — из последнего рассчитанного в памяти плана.
    """
    resolved_version_id = await _resolve_version_id(db, org_id, version_id)

    # Если есть сохранённая версия — грузим из БД
    if resolved_version_id is not None:
        col_check = await db.execute(
            text("""
                SELECT column_name FROM information_schema.columns
                WHERE table_name = 'scheduled_task'
                  AND column_name IN ('linked_equipment_id', 'task_role')
            """)
        )
        available_cols = {row.column_name for row in col_check.fetchall()}
        has_linked = "linked_equipment_id" in available_cols
        has_role = "task_role" in available_cols

        if has_linked and has_role:
            query = text("""
                SELECT
                    st.id, st.batch_id,
                    st.planned_start AS start, st.planned_end AS end,
                    eq.name AS equipment_name, eq.id AS equipment_id,
                    leq.name AS linked_equipment_name, leq.id AS linked_equipment_id,
                    p.name AS product_name, p.code AS product_code, p.id AS product_id,
                    ot.name AS operation_name,
                    st.task_role,
                    COALESCE(b.is_lab_blocked, FALSE) AS is_lab_blocked,
                    b.lab_status AS lab_status,
                    b.lab_block_reason AS lab_block_reason,
                    st.cooling_mode
                FROM scheduled_task st
                LEFT JOIN equipment eq ON st.equipment_id = eq.id
                LEFT JOIN equipment leq ON st.linked_equipment_id = leq.id
                LEFT JOIN operation_template ot ON st.operation_template_id = ot.id
                LEFT JOIN batch b ON st.batch_id = b.id
                LEFT JOIN product p ON b.product_id = p.id
                WHERE st.schedule_version_id = :version_id
                  AND st.organization_id = :org_id
                ORDER BY eq.name, st.planned_start ASC
            """)
        else:
            query = text("""
                SELECT
                    st.id, st.batch_id,
                    st.planned_start AS start, st.planned_end AS end,
                    eq.name AS equipment_name, eq.id AS equipment_id,
                    NULL::text AS linked_equipment_name, NULL::uuid AS linked_equipment_id,
                    p.name AS product_name, p.code AS product_code, p.id AS product_id,
                    ot.name AS operation_name,
                    NULL::text AS task_role,
                    COALESCE(b.is_lab_blocked, FALSE) AS is_lab_blocked,
                    b.lab_status AS lab_status,
                    b.lab_block_reason AS lab_block_reason,
                    st.cooling_mode
                FROM scheduled_task st
                LEFT JOIN equipment eq ON st.equipment_id = eq.id
                LEFT JOIN operation_template ot ON st.operation_template_id = ot.id
                LEFT JOIN batch b ON st.batch_id = b.id
                LEFT JOIN product p ON b.product_id = p.id
                WHERE st.schedule_version_id = :version_id
                  AND st.organization_id = :org_id
                ORDER BY eq.name, st.planned_start ASC
            """)

        result = await db.execute(query, {"version_id": resolved_version_id, "org_id": org_id})
        rows = result.fetchall()

        if not rows:
            raise HTTPException(
                status_code=404,
                detail="План с указанным version_id не найден или пуст",
            )

        gantt_tasks = []
        equipment_set = set()
        product_set = set()

        for row in rows:
            eq_name = str(row.equipment_name or row.equipment_id)[:30] if row.equipment_id else "Unknown"
            prod_name = str(row.product_name or row.product_code or row.product_id)[:30] if row.product_id else "Unknown"
            batch_name = f"Партия {str(row.batch_id)[:8]}" if row.batch_id else "Unknown"

            equipment_set.add(eq_name)
            if row.linked_equipment_name:
                equipment_set.add(str(row.linked_equipment_name)[:30])
            product_set.add(prod_name)

            duration_mins = int((row.end - row.start).total_seconds() / 60) if row.start and row.end else 0

            gantt_tasks.append(GanttTask(
                id=str(row.id),
                batch_id=batch_name,
                operation_name=str(row.operation_name or "Операция"),
                equipment_id=eq_name,
                product_id=prod_name,
                start=row.start,
                end=row.end,
                duration_minutes=duration_mins,
                linked_equipment_id=str(row.linked_equipment_id) if row.linked_equipment_id else None,
                linked_equipment_name=str(row.linked_equipment_name) if row.linked_equipment_name else None,
                task_role=row.task_role,
                is_lab_blocked=bool(row.is_lab_blocked) if row.is_lab_blocked is not None else False,
                lab_status=row.lab_status,
                lab_block_reason=row.lab_block_reason,
                cooling_mode=row.cooling_mode,   # Итерация 7
            ))

        makespan_hours = 0.0
        if gantt_tasks:
            max_end = max(task.end for task in gantt_tasks if task.end)
            min_start = min(task.start for task in gantt_tasks if task.start)
            makespan_hours = (max_end - min_start).total_seconds() / 3600

        return GanttResponse(
            tasks=gantt_tasks,
            total_tasks=len(gantt_tasks),
            makespan_hours=makespan_hours,
            equipment_list=sorted(list(equipment_set)),
            product_list=sorted(list(product_set)),
        )

    # === Fallback: из памяти ===
    if not _last_schedule_result:
        raise HTTPException(
            status_code=404,
            detail="Нет данных. Сначала постройте план на вкладке 'Планирование'",
        )

    tasks = _last_schedule_result.get("tasks", [])
    gantt_tasks = []
    equipment_set = set()
    product_set = set()

    for task in tasks:
        eq_name = str(task.get("equipment_name") or "")
        eq_id = str(task.get("equipment_id") or "")
        linked_eq_name = task.get("linked_equipment_name")
        linked_eq_id = task.get("linked_equipment_id")
        prod_name = str(task.get("product_name") or "")
        prod_id = str(task.get("product_id") or "")
        batch_name = str(task.get("batch_name") or "")
        batch_id = str(task.get("batch_id") or "")

        primary_label = eq_name if eq_name else eq_id[:20]
        equipment_set.add(primary_label)
        if linked_eq_name:
            equipment_set.add(linked_eq_name)
        product_set.add(prod_name if prod_name else prod_id[:20])

        start_dt = task.get("start")
        end_dt = task.get("end")

        if isinstance(start_dt, str):
            start_dt = datetime.fromisoformat(start_dt.replace('Z', '+00:00'))
        if isinstance(end_dt, str):
            end_dt = datetime.fromisoformat(end_dt.replace('Z', '+00:00'))

        duration_mins = int(task.get("duration", 0))
        if start_dt and end_dt and isinstance(start_dt, datetime) and isinstance(end_dt, datetime):
            duration_mins = int((end_dt - start_dt).total_seconds() / 60)

        gantt_tasks.append(GanttTask(
            id=f"{batch_id}_{task.get('op_id', '')}",
            batch_id=batch_name if batch_name else batch_id[:20],
            operation_name=str(task.get("operation_name", "Операция")),
            equipment_id=primary_label,
            product_id=prod_name if prod_name else prod_id[:20],
            start=start_dt,
            end=end_dt,
            duration_minutes=duration_mins,
            linked_equipment_id=str(linked_eq_id) if linked_eq_id else None,
            linked_equipment_name=linked_eq_name,
            task_role=task.get("role"),
            is_lab_blocked=False,
            lab_status=None,
            lab_block_reason=None,
            cooling_mode=task.get("cooling_mode"),   # Итерация 7
        ))

    return GanttResponse(
        tasks=gantt_tasks,
        total_tasks=len(gantt_tasks),
        makespan_hours=float(_last_schedule_result.get("makespan_minutes", 0)) / 60,
        equipment_list=sorted(list(equipment_set)),
        product_list=sorted(list(product_set)),
    )


@router.get("/export")
async def export_gantt_to_excel(
        version_id: UUID | None = Query(default=None),
        org_id: UUID = Depends(get_current_org_id),
        db: AsyncSession = Depends(get_db_session),
):
    """Экспорт Ганта в Excel."""
    response_data = await get_gantt_data(version_id=version_id, org_id=org_id, db=db)
    tasks = response_data.tasks

    wb = Workbook()
    ws_table = wb.active
    ws_table.title = "Расписание"

    headers = [
        "Оборудование", "Связанное оборудование", "Роль", "Операция",
        "Партия", "Продукт", "Начало", "Конец", "Длительность (мин)",
        "Заблокировано", "Причина", "Режим охлаждения",
    ]
    ws_table.append(headers)

    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws_table[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for task in tasks:
        start_str = task.start.strftime("%Y-%m-%d %H:%M") if isinstance(task.start, datetime) else str(task.start)
        end_str = task.end.strftime("%Y-%m-%d %H:%M") if isinstance(task.end, datetime) else str(task.end)

        ws_table.append([
            task.equipment_id,
            task.linked_equipment_name or "",
            task.task_role or "",
            task.operation_name,
            task.batch_id,
            task.product_id,
            start_str,
            end_str,
            task.duration_minutes,
            "Да" if task.is_lab_blocked else "Нет",
            task.lab_block_reason or "",
            task.cooling_mode or "",   # Итерация 7
        ])

    for col in ws_table.columns:
        max_len = max(len(str(cell.value)) for cell in col if cell.value is not None)
        ws_table.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    filename = f"APS_Gantt_{version_id if version_id else 'draft'}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )